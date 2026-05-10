"""文档解析器 — 支持多种文件格式"""

import os


def parse_document(file_path: str, content_type: str) -> str | None:
    """解析文档内容为纯文本"""
    if not file_path or not os.path.exists(file_path):
        return None

    ext = os.path.splitext(file_path)[1].lower()

    parsers = {
        ".md": _parse_markdown,
        ".txt": _parse_text,
        ".pdf": _parse_pdf,
        ".docx": _parse_docx,
        ".html": _parse_html,
        ".htm": _parse_html,
        ".xlsx": _parse_excel,
        ".xls": _parse_excel,
        ".csv": _parse_csv,
    }

    parser = parsers.get(ext)
    if parser:
        return parser(file_path)

    # 按 content_type 回退
    if "markdown" in content_type or "text" in content_type:
        return _parse_text(file_path)

    return None


def _parse_markdown(file_path: str) -> str:
    """解析 Markdown 文件"""
    with open(file_path, "r", encoding="utf-8", errors="replace") as f:
        return f.read()


def _parse_text(file_path: str) -> str:
    with open(file_path, "r", encoding="utf-8", errors="replace") as f:
        return f.read()


def _parse_pdf(file_path: str) -> str:
    """解析 PDF 文件（使用 PyPDF2）"""
    try:
        from PyPDF2 import PdfReader
        reader = PdfReader(file_path)
        text = []
        for page in reader.pages:
            text.append(page.extract_text() or "")
        return "\n\n".join(text)
    except ImportError:
        return f"[PDF 解析需要安装 PyPDF2] {os.path.basename(file_path)}"


def _parse_docx(file_path: str) -> str:
    """解析 DOCX 文件"""
    try:
        import docx
        doc = docx.Document(file_path)
        return "\n".join(p.text for p in doc.paragraphs)
    except ImportError:
        return f"[DOCX 解析需要安装 python-docx] {os.path.basename(file_path)}"


def _parse_html(file_path: str) -> str:
    """解析 HTML 文件"""
    try:
        import trafilatura
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            content = f.read()
        result = trafilatura.extract(content, output_format="txt")
        return result or content
    except ImportError:
        return f"[HTML 解析需要安装 trafilatura] {os.path.basename(file_path)}"


def _parse_excel(file_path: str) -> str:
    """解析 Excel 文件"""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        texts = []
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            rows_text = []
            for row in ws.iter_row():
                cells = [str(cell.value or "") for cell in row]
                rows_text.append(" | ".join(cells))
            texts.append(f"=== Sheet: {sheet} ===\n" + "\n".join(rows_text))
        return "\n\n".join(texts)
    except ImportError:
        return f"[Excel 解析需要安装 openpyxl] {os.path.basename(file_path)}"


def _parse_csv(file_path: str) -> str:
    """解析 CSV 文件"""
    try:
        import csv
        texts = []
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            reader = csv.reader(f)
            for row in reader:
                texts.append(" | ".join(row))
        return "\n".join(texts)
    except Exception as e:
        return f"[CSV 解析失败] {str(e)}"
